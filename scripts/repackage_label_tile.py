#!/usr/bin/env python3
"""Replay an existing verified label/tile manifest; copy bytes, never fetch or compress."""
import argparse
import hashlib
import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core.label_tile_manifest import assign_output_names, write_manifest
from openpyxl import Workbook, load_workbook
from PIL import Image


def repackage(source, destination):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if destination.exists() or destination == source or source in destination.parents:
        raise ValueError('输出目录必须是原包之外的全新目录')
    manifests = list(source.glob('*.xlsx'))
    if len(manifests) != 1:
        raise ValueError('必须有且仅有一个结果 Excel')
    values = list(load_workbook(manifests[0]).active.values)
    rows = [dict(zip(values[0], cells)) for cells in values[1:]]
    originals = []
    for row in rows:
        if row.get('下载结果') != '已下载':
            continue
        path = Path(str(row['本地文件'])).resolve()
        if source not in path.parents or not path.is_file():
            raise ValueError(f'源文件不在原包或不存在: {path}')
        if not str(row.get('云盘路径') or '') or not str(row.get('匹配策略') or ''):
            raise ValueError('缺少原素材来源或分类依据')
        if row.get('素材类型') == '鞋盒标签图/电子吊牌图' and 'OCR确认' not in str(row.get('备注') or ''):
            raise ValueError('标签缺少已确认的识别依据，不能根据 UUID 推断用途')
        with Image.open(path) as image:
            image.verify()
        originals.append((row, path, hashlib.sha256(path.read_bytes()).hexdigest()))
    assign_output_names(rows)
    destination.mkdir(parents=True, exist_ok=False)
    verification = []
    for row, original, digest in originals:
        target = destination / str(row['输入款号']) / row['输出文件名']
        target.parent.mkdir(exist_ok=True)
        if target.exists():
            raise ValueError(f'输出路径冲突: {target}')
        shutil.copy2(original, target)
        assert hashlib.sha256(target.read_bytes()).hexdigest() == digest
        assert hashlib.sha256(original.read_bytes()).hexdigest() == digest
        row['本地文件'] = str(target)
        verification.append({'输入款号': row['输入款号'], '原路径': str(original), '输出路径': str(target), 'sha256': digest})
    write_manifest(destination, rows)
    workbook = Workbook()
    keys = list(dict.fromkeys([*values[0], *(key for row in rows for key in row if not key.startswith('__'))]))
    workbook.active.title = '素材清单'
    workbook.active.append(keys)
    for row in rows:
        workbook.active.append([row.get(key) for key in keys])
    workbook.save(destination / '深绘吊牌洗唛平铺图下载结果.xlsx')
    (destination / '字节核验.json').write_text(json.dumps(verification, ensure_ascii=False, indent=2))
    return rows


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', required=True)
    parser.add_argument('--destination', required=True)
    args = parser.parse_args()
    rows = repackage(args.source, args.destination)
    print(json.dumps({'输出目录': args.destination, '已下载文件': sum(r.get('下载结果') == '已下载' for r in rows)}, ensure_ascii=False))
