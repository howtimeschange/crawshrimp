import json
from pathlib import Path

import pytest
from PIL import Image
from openpyxl import Workbook, load_workbook

from core.label_tile_manifest import assign_output_names
from scripts.repackage_label_tile import repackage


def asset(name, kind='款色图'):
    return {'输入款号': '2084261412A7', '输入编码': '2084261412A7', '文件名': name,
            '素材类型': kind, '云盘路径': f'冬/轻供/鞋品/2084261412A7/{name}',
            '匹配策略': '轻供 SKU 目录', '下载结果': '已下载', '本地文件': f'/tmp/{name}',
            '备注': 'OCR确认轻供电子合格证' if kind != '款色图' else ''}


def test_names_preserve_alphanumeric_style_and_original_without_color_inference():
    rows = [asset('31.jpg'), asset('32.jpg'), asset('1200-4.jpg'),
            asset('daf0269c-5d45-414f-aba3-f2f31bfa6dd3.png', '鞋盒标签图/电子吊牌图')]
    assign_output_names(rows)
    assert [r['输出文件名'] for r in rows] == [
        '2084261412A7_平铺图_01.jpg', '2084261412A7_平铺图_02.jpg',
        '2084261412A7_平铺图_03.jpg', '2084261412A7_吊牌_01.png']
    assert rows[0]['原文件名'] == '31.jpg'
    assert rows[3]['标签货号核验'].startswith('未核验')
    assert rows[3]['云盘目录归属'] == '完整款号目录匹配'
    assert not any('色码' in k or '成分' in k for r in rows for k in r)


def test_missing_material_is_not_renamed_as_success():
    row = {'输入款号': '2084261400U6', '下载结果': '未匹配到素材', '素材类型': '吊牌'}
    assign_output_names([row])
    assert row['输出文件名'] == ''
    assert row['输入款号'] == '2084261400U6'


def test_repackage_is_byte_preserving_and_updates_excel(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    row = asset('31.jpg')
    file = source / '31.jpg'
    Image.new('RGB', (20, 20), 'blue').save(file)
    original = file.read_bytes()
    row['本地文件'] = str(file)
    workbook = Workbook()
    workbook.active.append(list(row))
    workbook.active.append(list(row.values()))
    workbook.save(source / 'original.xlsx')
    out = tmp_path / 'out'
    result = repackage(source, out)
    target = Path(result[0]['本地文件'])
    assert target.read_bytes() == original == file.read_bytes()
    assert target.name == '2084261412A7_平铺图_01.jpg'
    manifest = json.loads((out / '素材清单.json').read_text())
    assert manifest['汇总']['实际款数'] == 1
    values = list(load_workbook(out / '深绘吊牌洗唛平铺图下载结果.xlsx').active.values)
    row = dict(zip(values[0], values[1]))
    assert row['文件名'] == row['输出文件名'] == target.name
    assert row['本地文件'] == str(target)
    with pytest.raises(ValueError):
        repackage(source, out)


def test_unknown_uuid_is_not_accepted_as_label(tmp_path):
    source = tmp_path / 'source'
    source.mkdir()
    row = asset('daf0269c-5d45-414f-aba3-f2f31bfa6dd3.png', '鞋盒标签图/电子吊牌图')
    row['备注'] = ''
    file = source / row['文件名']
    Image.new('RGB', (20, 20)).save(file)
    row['本地文件'] = str(file)
    w = Workbook()
    w.active.append(list(row)); w.active.append(list(row.values())); w.save(source / 'original.xlsx')
    with pytest.raises(ValueError, match='识别依据'):
        repackage(source, tmp_path / 'out')
    assert not (tmp_path / 'out').exists()


def test_status_directory_keeps_listingify_style_and_auditable_ocr_status():
    from core.label_tile_manifest import assign_output_names
    rows = [
        {'输入款号': '2081271412V1', '素材类型': kind, '下载结果': '已下载',
         '云盘路径': f'冬/轻供/鞋品/幻视/2081271412V1-已选/{name}',
         '本地文件': f'/tmp/{name}', '匹配策略': '白底文件名' if kind == '款色图' else 'OCR完整款号和标签结构',
         '标签货号核验': '不适用' if kind == '款色图' else 'OCR已核验完整款号；未推断业务色码'}
        for kind, name in [('款色图', '_S6A1141白底.jpg'), ('鞋盒标签图/电子吊牌图', '鞋盒.JPG')]
    ]
    assign_output_names(rows)
    assert [r['输出文件名'] for r in rows] == ['2081271412V1_平铺图_01.jpg', '2081271412V1_吊牌_01.jpg']
    assert all(r['云盘目录归属'] == '完整款号目录匹配' for r in rows)
    assert rows[1]['标签货号核验'].startswith('OCR已核验完整款号')
    assert rows[1]['原文件名'] == '鞋盒.JPG'
