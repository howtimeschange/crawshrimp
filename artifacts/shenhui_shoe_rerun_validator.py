#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core import shenhui_shoe_packaging as shoe  # noqa: E402


BASE_COLUMNS = [
    "输入款号",
    "颜色",
    "原文件名",
    "云盘路径",
    "规则槽位",
    "输出文件名",
    "处理动作",
    "下载结果",
    "本地文件",
    "压缩结果",
    "规则告警",
    "品类来源",
    "备注",
]


def text(value: Any) -> str:
    return str(value or "").strip()


def summarize_logs(logs: list[Any]) -> dict[str, Any]:
    lines = [text(item) for item in logs]
    correction_lines = [line for line in lines if "鞋品确定性校验：" in line]
    return {
        "correction_count": len(correction_lines),
        "pose_model_attempts": sum(
            "鞋品姿势识别模型" in line and "尝试：" in line
            for line in lines
        ),
        "ocr_model_attempts": sum("鞋盒标签 OCR 模型尝试" in line for line in lines),
        "soft_timeout_count": sum("60 秒软超时" in line for line in lines),
        "timeout_probe_count": sum("单批耐心复测" in line or "单次耐心复测" in line for line in lines),
        "fallback_count": sum("快速 fallback" in line for line in lines),
        "strategy_lines": [line for line in lines if "鞋品姿势识别策略：" in line],
    }


def rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [text(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        rows.append({
            headers[index]: (values[index] if index < len(values) else "")
            for index in range(len(headers))
        })
    wb.close()
    return rows


def prepared_rows(source_root: Path, report_xlsx: Path, style: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in rows_from_xlsx(report_xlsx):
        if text(row.get("输入款号")) != style or text(row.get("规则槽位")) != "原始素材":
            continue
        output_name = text(row.get("输出文件名"))
        if not output_name:
            continue
        local_path = source_root / style / output_name
        current = {key: ("" if value is None else value) for key, value in row.items()}
        current["本地文件"] = str(local_path)
        current["下载结果"] = "已下载"
        current["__shenhui_group_code"] = style
        current["__shoe_original_filename"] = text(row.get("原文件名")) or local_path.name
        color_folder = output_name.split("/", 1)[0]
        color_match = (
            re.search(r"(\d{5})", color_folder)
            or re.search(r"(\d{5})", text(row.get("颜色")))
        )
        current["__shoe_color_code"] = color_match.group(1) if color_match else ""
        rows.append(current)
    return rows


def write_report_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    extra_columns: list[str] = []
    for row in rows:
        for key in row:
            key_text = text(key)
            if key_text and key_text not in BASE_COLUMNS and not key_text.startswith("__"):
                extra_columns.append(key_text)
    columns = [*BASE_COLUMNS, *dict.fromkeys(extra_columns)]
    wb = Workbook()
    ws = wb.active
    ws.title = "结果"
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def source_exists_without_copy(report_rows: list[dict[str, Any]], color: str, source_name: str) -> bool:
    if "拷贝" not in source_name:
        return False
    source_key = shoe._copy_variant_key(source_name)
    for row in report_rows:
        if text(row.get("颜色")) != color or text(row.get("规则槽位")) != "原始素材":
            continue
        original_name = text(row.get("原文件名"))
        if "拷贝" not in original_name and shoe._copy_variant_key(original_name) == source_key:
            return True
    return False


def validate_style(
    *,
    style: str,
    style_root: Path,
    report_rows: list[dict[str, Any]],
    category: str,
) -> tuple[list[str], list[str]]:
    issues: list[str] = []
    warnings: list[str] = []
    root_required = [f"tmz ({index}).jpg" for index in range(1, 6)]
    root_required.extend(["tmq.jpg", "tmt.png"])
    for name in root_required:
        if not (style_root / name).is_file():
            issues.append(f"missing root {name}")
    color_dirs = sorted(
        [
            path
            for path in style_root.iterdir()
            if path.is_dir() and re.match(r"^\d+\.", path.name)
        ],
        key=lambda path: path.name,
    ) if style_root.is_dir() else []
    if not color_dirs:
        issues.append("missing color directories")
    if len(list(style_root.glob("jdt.*.png"))) < len(color_dirs):
        issues.append("missing root jdt.*.png for one or more colors")
    if len(list(style_root.glob("wpt30.*.png"))) < len(color_dirs):
        issues.append("missing root wpt30.*.png for one or more colors")

    for color_dir in color_dirs:
        required = [
            "wpz (1).jpg",
            "wpz (2).jpg",
            "wpz (3).jpg",
            "wpz (4).jpg",
            "wpz (15).jpg",
            "wpz (16).jpg",
            "tms.jpg",
        ]
        for name in required:
            if not (color_dir / name).is_file():
                issues.append(f"{color_dir.name} missing {name}")
        if not list(color_dir.glob("yq*.jpg")):
            issues.append(f"{color_dir.name} missing yq*.jpg")
        wpz15 = color_dir / "wpz (15).jpg"
        if wpz15.is_file():
            feature = shoe._binary_pose_feature(wpz15)
            if feature.background_luma >= shoe.SHOE_WHITE_BACKGROUND_LUMA:
                warnings.append(
                    f"{color_dir.name} wpz (15).jpg is not gray background; "
                    f"accepted as fallback when no qualified gray source exists "
                    f"luma={feature.background_luma:.1f}"
                )
    if color_dirs and not (color_dirs[0] / "o.jpg").is_file():
        issues.append(f"{color_dirs[0].name} missing o.jpg")

    tmz_features: list[tuple[int, shoe._BinaryPoseFeature]] = []
    for index in range(1, 6):
        path = style_root / f"tmz ({index}).jpg"
        if path.is_file():
            tmz_features.append((index, shoe._binary_pose_feature(path)))
    if len(tmz_features) == 5:
        category_text = text(category)
        pose3_max_aspect = 0.95 if category_text == "婴童" else 0.82
        pose3_max_coverage = 0.145 if category_text == "婴童" else 0.16
        pose3 = tmz_features[2][1]
        if not (
            0.45 <= pose3.aspect_ratio <= pose3_max_aspect
            and pose3.bounding_coverage <= pose3_max_coverage
        ):
            issues.append(
                "tmz (3).jpg invalid pose3 feature "
                f"aspect={pose3.aspect_ratio:.3f} coverage={pose3.bounding_coverage:.3f}"
            )
        pose5 = tmz_features[4][1]
        if pose5.background_luma < shoe.SHOE_WHITE_BACKGROUND_LUMA:
            issues.append(f"tmz (5).jpg is not white background luma={pose5.background_luma:.1f}")
        for first_index, first_feature in tmz_features:
            for second_index, second_feature in tmz_features:
                if first_index >= second_index:
                    continue
                distance = shoe._binary_pose_distance(first_feature, second_feature)
                if distance <= shoe.SHOE_MAIN_SLOT_DUPLICATE_MAX_DISTANCE:
                    issues.append(
                        f"tmz duplicate pose {first_index}/{second_index} "
                        f"distance={distance:.4f}"
                    )

    for row in report_rows:
        slot = text(row.get("规则槽位"))
        warning = text(row.get("规则告警"))
        action = text(row.get("处理动作"))
        if "已跳过 tmq" in warning or "已跳过 tmq" in action:
            issues.append("report contains skipped tmq")
        if slot in {
            "tmz1",
            "tmz2",
            "tmz3",
            "tmz4",
            "wpz1",
            "wpz2",
            "wpz3",
            "wpz4",
        }:
            color = text(row.get("颜色"))
            source_name = text(row.get("原文件名"))
            if source_exists_without_copy(report_rows, color, source_name):
                issues.append(f"{slot} uses copied white source {source_name}")
    return issues, warnings


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("style")
    parser.add_argument("--attempt", required=True)
    parser.add_argument("--source-root", default="/Users/xingyicheng/Downloads/鞋品测试")
    parser.add_argument(
        "--report-xlsx",
        default="/Users/xingyicheng/Downloads/鞋品测试/深绘鞋品上新图包整理结果_20260822-202441.xlsx",
    )
    parser.add_argument(
        "--category-xlsx",
        default="/Users/xingyicheng/Downloads/鞋品品类映射模板测试.xlsx",
    )
    parser.add_argument(
        "--output-root",
        default="/Users/xingyicheng/Downloads/鞋品测试-最终重跑校验-20260824-Sol链",
    )
    parser.add_argument(
        "--pose-models",
        default="gpt-5.6-sol,gpt-5.6-terra,gpt-5.6-luna,gpt-5.5",
        help="Comma-separated model chain for pose recognition; first is primary.",
    )
    parser.add_argument(
        "--label-models",
        default="gpt-5.6-sol,gpt-5.6-terra,gpt-5.6-luna,gpt-5.5",
        help="Comma-separated model chain for OCR; first is primary.",
    )
    parser.add_argument("--pose-parallelism", type=int, default=shoe.SHOE_POSE_BATCH_PARALLELISM)
    parser.add_argument("--pose-timeout", type=float, default=shoe.SHOE_POSE_MODEL_TIMEOUT_SECONDS)
    parser.add_argument(
        "--pose-strategy",
        default=shoe.SHOE_POSE_DEFAULT_STRATEGY,
        choices=[
            shoe.SHOE_POSE_STRATEGY_GLOBAL_PAGES,
            shoe.SHOE_POSE_STRATEGY_BATCH,
            shoe.SHOE_POSE_STRATEGY_BATCH_OVERVIEW,
            shoe.SHOE_POSE_STRATEGY_SINGLE_SHEET,
        ],
    )
    parser.add_argument("--label-timeout", type=float, default=shoe.SHOE_LABEL_OCR_TIMEOUT_SECONDS)
    parser.add_argument("--copy-final", action="store_true")
    args = parser.parse_args()

    style = text(args.style)
    output_root = Path(args.output_root)
    attempt_root = output_root / "_attempts" / style / f"attempt-{args.attempt}"
    if attempt_root.exists():
        shutil.rmtree(attempt_root)
    attempt_root.mkdir(parents=True, exist_ok=True)

    source_root = Path(args.source_root)
    category_rows = rows_from_xlsx(Path(args.category_xlsx))
    shoe_categories = shoe.parse_shoe_category_rows(category_rows)
    data_rows = prepared_rows(source_root, Path(args.report_xlsx), style)
    logs: list[str] = []
    pose_models = [text(item) for item in args.pose_models.split(",") if text(item)]
    label_models = [text(item) for item in args.label_models.split(",") if text(item)]
    if not pose_models:
        pose_models = ["gpt-5.6-sol"]
    if not label_models:
        label_models = ["gpt-5.6-sol"]
    shoe.SHOE_POSE_BATCH_PARALLELISM = max(1, int(args.pose_parallelism))
    shoe.SHOE_POSE_MODEL_TIMEOUT_SECONDS = max(1.0, float(args.pose_timeout))
    shoe.SHOE_LABEL_OCR_TIMEOUT_SECONDS = max(1.0, float(args.label_timeout))

    start = time.time()
    try:
        report_rows, package_roots = shoe.prepare_shoe_packages(
            data_rows=data_rows,
            output_root=attempt_root,
            model_id=pose_models[0],
            pose_strategy=args.pose_strategy,
            fallback_model_ids=pose_models[1:],
            label_model_id=label_models[0],
            label_fallback_model_ids=label_models[1:],
            shoe_categories=shoe_categories,
            log=logs.append,
        )
    except Exception as exc:
        elapsed = round(time.time() - start, 2)
        summary = {
            "style": style,
            "attempt": args.attempt,
            "elapsed": elapsed,
            "category": shoe_categories.get(style, ""),
            "pose_models": pose_models,
            "pose_strategy": args.pose_strategy,
            "label_models": label_models,
            "pose_parallelism": shoe.SHOE_POSE_BATCH_PARALLELISM,
            "pose_timeout": shoe.SHOE_POSE_MODEL_TIMEOUT_SECONDS,
            "label_timeout": shoe.SHOE_LABEL_OCR_TIMEOUT_SECONDS,
            "issues": [f"{type(exc).__name__}: {text(exc)}"],
            "metrics": summarize_logs(logs),
            "logs_tail": logs[-80:],
        }
        (attempt_root / "logs.txt").write_text("\n".join(logs), encoding="utf-8")
        (attempt_root / "validation.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 1
    elapsed = round(time.time() - start, 2)
    style_root = package_roots.get(style, attempt_root / style)
    report_json = attempt_root / "report_rows.json"
    report_json.write_text(json.dumps(report_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    write_report_xlsx(attempt_root / "report.xlsx", report_rows)
    (attempt_root / "logs.txt").write_text("\n".join(logs), encoding="utf-8")

    issues, validation_warnings = validate_style(
        style=style,
        style_root=style_root,
        report_rows=report_rows,
        category=shoe_categories.get(style, ""),
    )
    color_dirs = sorted(
        [
            path.name
            for path in style_root.iterdir()
            if path.is_dir() and re.match(r"^\d+\.", path.name)
        ],
    ) if style_root.is_dir() else []
    summary = {
        "style": style,
        "attempt": args.attempt,
        "elapsed": elapsed,
        "style_root": str(style_root),
        "report_rows": len(report_rows),
        "report_xlsx": str(attempt_root / "report.xlsx"),
        "color_dirs": color_dirs,
        "category": shoe_categories.get(style, ""),
        "pose_models": pose_models,
        "pose_strategy": args.pose_strategy,
        "label_models": label_models,
        "pose_parallelism": shoe.SHOE_POSE_BATCH_PARALLELISM,
        "pose_timeout": shoe.SHOE_POSE_MODEL_TIMEOUT_SECONDS,
        "label_timeout": shoe.SHOE_LABEL_OCR_TIMEOUT_SECONDS,
        "issues": issues,
        "warnings": validation_warnings,
        "metrics": summarize_logs(logs),
        "logs_tail": logs[-50:],
    }
    if not issues and args.copy_final:
        final_root = output_root / "final"
        final_style_root = final_root / style
        if final_style_root.exists():
            shutil.rmtree(final_style_root)
        final_root.mkdir(parents=True, exist_ok=True)
        shutil.copytree(style_root, final_style_root)
        shutil.copy2(attempt_root / "report.xlsx", final_root / f"{style}-report.xlsx")
        summary["final_style_root"] = str(final_style_root)
    (attempt_root / "validation.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not issues else 2


if __name__ == "__main__":
    raise SystemExit(main())
