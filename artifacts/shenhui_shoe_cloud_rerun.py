#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import time
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core import shenhui_shoe_packaging as shoe  # noqa: E402
from artifacts import shenhui_shoe_rerun_validator as validator  # noqa: E402


DEFAULT_CLOUD_PATH = (
    "巴拉营运BU-商品//巴拉货控/02 产品上新模块/2-2 巴拉产品上新/"
    "2026年巴拉秋/平拍原图/全域/小程序/鞋品/"
)
DEFAULT_STYLES = [
    "204426146036",
    "204426146127",
    "204426146023",
    "204426141113",
    "204426141112",
    "204426141127",
    "204426140034",
    "204426140143",
]
DEFAULT_CHAIN = [
    "gpt-5.6-sol",
    "gpt-5.6-terra",
    "gpt-5.6-luna",
    "gpt-5.5",
    "deepseek-official-v4-flash-vision-exp",
    "kimi-k2.7-code",
]
DEFAULT_LABEL_CHAIN = [
    "gpt-5.6-sol",
    "gemini-3.5-flash",
    "qwen3.7-plus",
    "gpt-5.6-terra",
    "kimi-k2.7-code",
    "deepseek-official-v4-flash-vision-exp",
]
TERMINAL_RUN_STATUSES = {
    "success",
    "done",
    "completed",
    "failed",
    "error",
    "cancelled",
    "stopped",
}
SUCCESS_RUN_STATUSES = {"success", "done", "completed"}


def text(value: Any) -> str:
    return str(value or "").strip()


def summarize_logs(logs: list[Any]) -> dict[str, Any]:
    lines = [text(item) for item in logs]
    correction_lines = [line for line in lines if "鞋品确定性校验：" in line]
    correction_slots = {
        "tmz1": 0,
        "tmz2": 0,
        "tmz3": 0,
        "tmz4": 0,
        "tmz5_wpz5": 0,
        "yq": 0,
        "other": 0,
    }
    for line in correction_lines:
        matched = False
        for key, patterns in {
            "tmz1": ("主图1", "tmz1", "wpz1"),
            "tmz2": ("主图2", "tmz2", "wpz2"),
            "tmz3": ("主图3", "第3姿势", "tmz3", "wpz3"),
            "tmz4": ("主图4", "tmz4", "wpz4"),
            "tmz5_wpz5": ("主图5", "tmz5", "wpz5"),
            "yq": ("yq1", "yq2", "yq3"),
        }.items():
            if any(pattern in line for pattern in patterns):
                correction_slots[key] += 1
                matched = True
        if not matched:
            correction_slots["other"] += 1
    model_attempt_lines = [line for line in lines if "鞋品姿势识别模型" in line and "尝试：" in line]
    ocr_attempt_lines = [line for line in lines if "鞋盒标签 OCR 模型尝试" in line]
    return {
        "correction_count": len(correction_lines),
        "correction_slots": correction_slots,
        "pose_model_attempts": len(model_attempt_lines),
        "ocr_model_attempts": len(ocr_attempt_lines),
        "soft_timeout_count": sum("60 秒软超时" in line for line in lines),
        "timeout_probe_count": sum("单批耐心复测" in line or "单次耐心复测" in line for line in lines),
        "fallback_count": sum("快速 fallback" in line for line in lines),
        "strategy_lines": [line for line in lines if "鞋品姿势识别策略：" in line],
    }


def request_json(
    base_url: str,
    token: str,
    method: str,
    path: str,
    payload: dict[str, Any] | None = None,
    timeout: float = 30,
) -> dict[str, Any]:
    body = None
    headers = {
        "X-Crawshrimp-Token": token,
        "Accept": "application/json",
    }
    if payload is not None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    request = Request(
        f"{base_url.rstrip('/')}/{path.lstrip('/')}",
        data=body,
        headers=headers,
        method=method,
    )
    try:
        with urlopen(request, timeout=timeout) as response:
            raw = response.read().decode("utf-8", errors="replace")
    except HTTPError as exc:
        raw = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code} {method} {path}: {raw[:500]}") from exc
    except (URLError, TimeoutError, OSError) as exc:
        raise RuntimeError(f"{method} {path} failed: {exc}") from exc
    if not raw.strip():
        return {}
    parsed = json.loads(raw)
    return parsed if isinstance(parsed, dict) else {"value": parsed}


def category_map(path: Path) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {}
    headers = [text(item) for item in rows[0]]
    result: dict[str, str] = {}
    for values in rows[1:]:
        row = {
            headers[index]: values[index] if index < len(values) else ""
            for index in range(len(headers))
        }
        style = text(row.get("款号"))
        if re.fullmatch(r"\d+\.0", style):
            style = style[:-2]
        category = text(row.get("品类"))
        if style:
            result[style] = category
    return result


def normalize_category(value: str) -> str:
    return shoe.normalize_shoe_category(value) or text(value)


def parse_output_files(value: Any) -> list[str]:
    if isinstance(value, list):
        return [text(item) for item in value if text(item)]
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return [value] if value.strip() else []
        return parse_output_files(parsed)
    return []


def latest_run_id(status: dict[str, Any]) -> int:
    last = status.get("last_run")
    if isinstance(last, dict):
        try:
            return int(last.get("id") or 0)
        except (TypeError, ValueError):
            return 0
    return 0


def run_one(
    *,
    base_url: str,
    token: str,
    style: str,
    category: str,
    export_root: Path,
    cloud_path: str,
    model_chain: list[str],
    label_model_chain: list[str],
    pose_strategy: str,
    timeout_seconds: float,
    poll_seconds: float,
    artifact_root: Path,
) -> dict[str, Any]:
    export_root.mkdir(parents=True, exist_ok=True)
    artifact_root.mkdir(parents=True, exist_ok=True)
    request_json(
        base_url,
        token,
        "DELETE",
        "/tasks/shenhui-new-arrival/prepare_shoe_upload_package/logs",
        timeout=10,
    )
    before = request_json(
        base_url,
        token,
        "GET",
        "/tasks/shenhui-new-arrival/prepare_shoe_upload_package/status",
        timeout=10,
    )
    previous_run_id = latest_run_id(before)
    params: dict[str, Any] = {
        "mode": "new",
        "shoe_cloud_path": cloud_path,
        "shoe_category_file": {
            "rows": [{"款号": style, "品类": category}],
        },
        "model_id": model_chain[0],
        "shoe_pose_strategy": pose_strategy,
        "package_name": style,
        "export_folder": str(export_root),
    }
    for index, model_id in enumerate(model_chain[1:6], start=1):
        params[f"fallback_model_{index}"] = model_id
    params["label_model_id"] = label_model_chain[0]
    params["label_fallback_model_ids"] = label_model_chain[1:6]

    started = time.time()
    request_json(
        base_url,
        token,
        "POST",
        "/tasks/shenhui-new-arrival/prepare_shoe_upload_package/run",
        {"params": params},
        timeout=30,
    )

    status: dict[str, Any] = {}
    last_log_count = 0
    while True:
        status = request_json(
            base_url,
            token,
            "GET",
            "/tasks/shenhui-new-arrival/prepare_shoe_upload_package/status",
            timeout=10,
        )
        logs_payload = request_json(
            base_url,
            token,
            "GET",
            "/tasks/shenhui-new-arrival/prepare_shoe_upload_package/logs",
            timeout=10,
        )
        logs = logs_payload.get("logs") if isinstance(logs_payload.get("logs"), list) else []
        if len(logs) != last_log_count:
            last_log_count = len(logs)
            (artifact_root / "logs.txt").write_text("\n".join(map(text, logs)), encoding="utf-8")
        live = status.get("live")
        last = status.get("last_run") if isinstance(status.get("last_run"), dict) else {}
        current_run_id = latest_run_id(status)
        if not live and current_run_id and current_run_id != previous_run_id:
            break
        live_status = text(live.get("status")) if isinstance(live, dict) else ""
        if live_status in TERMINAL_RUN_STATUSES and current_run_id != previous_run_id:
            break
        if time.time() - started > timeout_seconds:
            raise RuntimeError(f"{style} run timed out after {timeout_seconds:g}s")
        phase = text(live.get("phase")) if isinstance(live, dict) else ""
        records = live.get("records") if isinstance(live, dict) else ""
        print(
            json.dumps(
                {
                    "style": style,
                    "elapsed": round(time.time() - started, 1),
                    "status": live_status or text(last.get("status")),
                    "phase": phase,
                    "records": records,
                    "logs": len(logs),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        time.sleep(poll_seconds)

    logs_payload = request_json(
        base_url,
        token,
        "GET",
        "/tasks/shenhui-new-arrival/prepare_shoe_upload_package/logs",
        timeout=10,
    )
    logs = logs_payload.get("logs") if isinstance(logs_payload.get("logs"), list) else []
    last = status.get("last_run") if isinstance(status.get("last_run"), dict) else {}
    output_files = parse_output_files(last.get("output_files"))
    style_root = next(
        (
            Path(item).expanduser()
            for item in output_files
            if Path(item).expanduser().is_dir()
            and Path(item).expanduser().name == style
        ),
        export_root / style,
    )
    report_path = next(
        (
            Path(item).expanduser()
            for item in output_files
            if Path(item).expanduser().is_file()
            and Path(item).expanduser().suffix.lower() in {".xlsx", ".xlsm"}
        ),
        None,
    )
    report_rows = validator.rows_from_xlsx(report_path) if report_path and report_path.is_file() else []
    report_style_roots: list[Path] = []
    for row in report_rows:
        local_path = Path(text(row.get("本地文件"))).expanduser()
        for parent in local_path.parents:
            if re.fullmatch(rf"{re.escape(style)}(?:_\d+)?", parent.name):
                report_style_roots.append(parent)
                break
    if report_style_roots:
        style_root = report_style_roots[0]
    validation_result = validator.validate_style(
        style=style,
        style_root=style_root,
        report_rows=report_rows,
        category=normalize_category(category),
    )
    if isinstance(validation_result, tuple):
        issues, validation_warnings = validation_result
    else:
        issues = validation_result
        validation_warnings = []
    summary = {
        "style": style,
        "category": category,
        "normalized_category": normalize_category(category),
        "status": text(last.get("status")),
        "run_id": last.get("id"),
        "elapsed": round(time.time() - started, 2),
        "model_chain": model_chain,
        "label_model_chain": label_model_chain,
        "pose_strategy": pose_strategy,
        "export_root": str(export_root),
        "style_root": str(style_root),
        "report_path": str(report_path) if report_path else "",
        "output_files": output_files,
        "issues": issues,
        "warnings": validation_warnings,
        "metrics": summarize_logs(logs),
        "logs_tail": [text(item) for item in logs[-80:]],
    }
    (artifact_root / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    if text(last.get("status")) not in SUCCESS_RUN_STATUSES:
        raise RuntimeError(f"{style} task status is {last.get('status')}: {last.get('error')}")
    if issues:
        raise RuntimeError(f"{style} validation failed: {'; '.join(issues[:8])}")
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:18766")
    parser.add_argument("--token-file", default="/tmp/crawshrimp-codex-shenhui-lock/api-token")
    parser.add_argument("--category-xlsx", default="/Users/xingyicheng/Downloads/鞋品品类映射模板测试.xlsx")
    parser.add_argument("--cloud-path", default=DEFAULT_CLOUD_PATH)
    parser.add_argument("--output-root", default="/Users/xingyicheng/Downloads/鞋品干净云盘重跑-20260825-Kimi链")
    parser.add_argument("--artifact-root", default="artifacts/shenhui-shoe-cloud-rerun-20260825")
    parser.add_argument("--styles", default=",".join(DEFAULT_STYLES))
    parser.add_argument("--model-chain", default=",".join(DEFAULT_CHAIN))
    parser.add_argument(
        "--label-model-chain",
        default=",".join(DEFAULT_LABEL_CHAIN),
        help="comma-separated cross-family label OCR routes",
    )
    parser.add_argument(
        "--pose-strategy",
        default=shoe.SHOE_POSE_DEFAULT_STRATEGY,
        choices=[
            shoe.SHOE_POSE_STRATEGY_GLOBAL_PAGES,
            shoe.SHOE_POSE_STRATEGY_BATCH,
            shoe.SHOE_POSE_STRATEGY_BATCH_OVERVIEW,
            shoe.SHOE_POSE_STRATEGY_SINGLE_SHEET,
        ],
    )
    parser.add_argument("--timeout", type=float, default=3600)
    parser.add_argument("--poll", type=float, default=10)
    parser.add_argument("--clean-output", action="store_true")
    args = parser.parse_args()

    token = Path(args.token_file).expanduser().read_text(encoding="utf-8").strip()
    categories = category_map(Path(args.category_xlsx).expanduser())
    styles = [text(item) for item in re.split(r"[,，;；\s]+", args.styles) if text(item)]
    model_chain = [text(item) for item in re.split(r"[,，;；\s]+", args.model_chain) if text(item)]
    if not model_chain:
        raise SystemExit("model chain is empty")
    label_model_chain = [
        text(item)
        for item in re.split(r"[,，;；\s]+", args.label_model_chain)
        if text(item)
    ] or list(DEFAULT_LABEL_CHAIN)
    output_root = Path(args.output_root).expanduser()
    artifact_root = Path(args.artifact_root)
    if args.clean_output and output_root.exists():
        shutil.rmtree(output_root)
    output_root.mkdir(parents=True, exist_ok=True)
    artifact_root.mkdir(parents=True, exist_ok=True)

    results = []
    for style in styles:
        category = categories.get(style, "")
        if not category:
            raise RuntimeError(f"missing category for {style} in {args.category_xlsx}")
        summary = run_one(
            base_url=args.base_url,
            token=token,
            style=style,
            category=category,
            export_root=output_root,
            cloud_path=args.cloud_path,
            model_chain=model_chain,
            label_model_chain=label_model_chain,
            pose_strategy=args.pose_strategy,
            timeout_seconds=args.timeout,
            poll_seconds=args.poll,
            artifact_root=artifact_root / style,
        )
        results.append(summary)
    all_summary = {
        "output_root": str(output_root),
        "artifact_root": str(artifact_root),
        "model_chain": model_chain,
        "label_model_chain": label_model_chain,
        "pose_strategy": args.pose_strategy,
        "results": results,
    }
    (artifact_root / "summary.json").write_text(
        json.dumps(all_summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(all_summary, ensure_ascii=False, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
